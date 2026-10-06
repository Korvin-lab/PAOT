# -*- coding: utf-8 -*-
"""Publish the verified weekly file and compact Russian source/coverage reports."""
import csv
import json
import math
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
FINAL = HERE / "final_dataset__WEEKLY_7D.csv"
SOURCE_SHA = "5d05ff93c7450aa166be03f2f9e89c50c85dfbd7a87c09b75e755887e91c1d7b"


def main():
    verification = json.loads((HERE / "independent_weekly_verification.json").read_text())
    assert verification["status"] == "PASS" and verification["columns"] == 50
    plan = json.loads((HERE / "plan_report.json").read_text())
    h2s = json.loads((HERE / "h2s_source_audit.json").read_text())
    partial = HERE / "final_dataset__WEEKLY_7D.partial.csv"
    if partial.exists():
        assert not FINAL.exists()
        partial.rename(FINAL)
    assert FINAL.exists()
    n = verification["weekly_rows"]
    full = plan["days_per_week"]["7"]
    incomplete = verification["pipe_weeks"] - full
    with (HERE / "Сводка_по_трубам_7_дней.csv").open(encoding="utf-8-sig") as f:
        pipe_rows = list(csv.DictReader(f))
    assert len(pipe_rows) == 66
    assert sum(int(r["Рабочих дней в исходном CSV"]) for r in pipe_rows) == 71460
    manifest = {"status": "PASS", "source": str(HERE.parent / "megion_kvch_ing_rebuild_2026-10-01/final_dataset__KVCH_ING50.csv"),
                "source_sha256": SOURCE_SHA, "source_rows": 138583872, "output": str(FINAL),
                "output_sha256": verification["sha256"], "output_bytes": FINAL.stat().st_size,
                "output_rows": n, "columns": 50, "pipes": 66, "source_pipe_dates": 71460,
                "pipe_weeks": verification["pipe_weeks"], "complete_7_day_pipe_weeks": full,
                "partial_pipe_weeks": incomplete, "first_week": plan["first_week"], "last_week": plan["last_week"],
                "weekly_rule": "Monday-Sunday, date labels the Monday; group by id and segment_id; arithmetic mean of available finite daily values; all-NaN remains NaN",
                "identifiers": "id and segment_id preserved, never averaged; no rows created for entirely unobserved weeks",
                "h2s_oil_unique_target_measurements": 15, "h2s_oil_positive_target_measurements": 0,
                "h2s_water_policy": "Preserve missing values in the existing daily dataset; source oil zeros are audited and not silently converted into positive concentrations"}
    (HERE / "FINAL_MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    workbook = Workbook(); summary = workbook.active; summary.title = "Итог"
    coverage = workbook.create_sheet("По трубам"); evidence = workbook.create_sheet("H2S в исходниках")
    summary_rows = [
        ["Показатель", "Значение / пояснение"],
        ["Недельный файл", FINAL.name], ["Недельных сегментных строк", n], ["Колонок", 50],
        ["Труб", 66], ["Исходных трубо-дней", 71460], ["Трубо-недель", verification["pipe_weeks"]],
        ["Полных недель (7 дней)", full], ["Неполных недель", incomplete],
        ["Сохранение исходного покрытия", "Использованы все исходные трубо-дни и сегменты; новых замеров или отсутствующих недель не добавлено."],
        ["Правило дат", "Понедельник–воскресенье; date обозначает понедельник. Первые/последние недели могут быть неполными."],
        ["Среднее", "По каждому ID трубы и segment_id отдельно; по имеющимся непустым значениям. Пустая целиком неделя параметра остается пустой."],
        ["Нулевые значения", "Участвуют в средних там, где присутствуют в исходном датасете; пропуски нулями не заменяются."],
        ["Производные параметры", "Усредняются готовые значения каждого параметра. Формулы по средним входам и гидравлика повторно не рассчитывались."],
        ["Проверено УКК", "128 файлов, 228 листов, обе папки УКК от ДО / УКК от ДО 2; шапки и заполненные H2S-ячейки перечитаны."],
        ["Нефтяной H2S для наших труб", "30 исходных ячеек, 15 уникальных сочетаний УКК–дата–значение после удаления повторов; все равны 0 мг/м3. Через 4 направления относятся к 12 трубам."],
        ["Положительный нефтяной H2S", "3,7 мг/м3, УКК 604, 28.04.2026. Направление 1750009316 отсутствует среди наших 66 труб. Запись повторяется в двух книгах."],
        ["Почему H2S воды пуст", "Действующая сборка допускает только положительные H2S при заполнении. Исходные нефтяные нули были исключены. Их лабораторный смысл отдельно не подтвержден."],
        ["Правило фаз", "Нефтяная фаза принимается водной по правилу пользователя; мг/м3 делятся на 1000 для мг/л. Газ в воду не переносится."],
        ["Критично", "Водный H2S пуст; происхождение КВЧ и условный факт/план ing_factor сохраняют ограничения ежедневного датасета."],
    ]
    for row in summary_rows: summary.append(row)
    summary.column_dimensions["A"].width = 39; summary.column_dimensions["B"].width = 105
    for row in summary.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row[0].row].height = 18*max(1, math.ceil(len(str(row[1].value))/85))
    headers = list(pipe_rows[0]); coverage.append(headers)
    for row in pipe_rows:
        coverage.append([row[c] if c == "ID простого участка" else int(row[c]) for c in headers])
    for c in range(1, len(headers)+1):
        coverage.column_dimensions[coverage.cell(1,c).column_letter].width = 23 if c == 1 else 28
    coverage.row_dimensions[1].height = 48
    source_rows = h2s["target_liquid"] + h2s["positive_liquid"]
    unique = {}
    for row in source_rows:
        key = (row["Номер УКК"], row["Дата"], row["Числовое значение"], row["Единица"])
        unique.setdefault(key, row)
    evidence.append(["Номер УКК", "Дата", "H2S в нефти, мг/м3", "Направление по ПОТ", "Наши трубы", "Файл", "Лист", "Ячейка"])
    evidence.row_dimensions[1].height = 32
    for _, row in sorted(unique.items(), key=lambda item: (item[0][1] or "", item[0][0])):
        evidence.append([row["Номер УКК"], datetime.fromisoformat(row["Дата"]), row["Числовое значение"],
                         row["Направления ПОТ"], row["Целевые трубы"] or "Вне выборки 66 труб", row["Файл"], row["Лист"], row["Ячейка"]])
    for row in evidence.iter_rows(min_row=2):
        row[1].number_format = "dd.mm.yyyy"
        for cell in row: cell.alignment = Alignment(wrap_text=True, vertical="top")
        evidence.row_dimensions[row[0].row].height = 42
    for col, width in {"A":14,"B":15,"C":23,"D":24,"E":57,"F":85,"G":18,"H":15}.items():
        evidence.column_dimensions[col].width = width
    for ws in workbook:
        for cell in ws[1]:
            cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True, vertical="center")
    path = HERE / "Мегион_недели_и_проверка_H2S_2026_10_02.xlsx"
    workbook.save(path)
    with zipfile.ZipFile(path) as z: assert z.testzip() is None
    checked = load_workbook(path, read_only=True, data_only=True)
    assert checked["По трубам"].max_row == 67 and checked["H2S в исходниках"].max_row == 17
    checked.close()

    text = f"""# Мегион: недельные средние и повторная проверка H2S (02.10.2026)

Готовый файл: `final_dataset__WEEKLY_7D.csv`, {n:,} строк, 50 колонок, 66 труб, {FINAL.stat().st_size:,} байт. SHA-256: `{verification['sha256']}`. Исходный ежедневный CSV сохранен: `../megion_kvch_ing_rebuild_2026-10-01/final_dataset__KVCH_ING50.csv`.

## Как получены недельные значения

Неделя — понедельник–воскресенье, дата строки обозначает понедельник. Период меток недель: {plan['first_week']}..{plan['last_week']}. ID трубы и segment_id сохранены. Остальные 47 числовых полей усреднены по той же трубе и сегменту, отдельно по непустым значениям каждого поля. Нуль в готовом ежедневном CSV участвует в среднем; NaN не участвует; все-NaN остаются NaN. При отсутствии всей недели новые строки не создаются. Средние вычисляются по имеющимся дням, а не делением суммы всегда на 7.

Все 138 583 872 исходные строки и 71 460 трубо-дней использованы ровно один раз. Получено 11 357 трубо-недель: {full} содержат 7 дней, {incomplete} неполные. Подробная сводка для 66 труб — `Сводка_по_трубам_7_дней.csv` и лист `По трубам` книги `Мегион_недели_и_проверка_H2S_2026_10_02.xlsx`. Готовые производные величины тоже усредняются; формулы по средним входам и гидравлика заново не рассчитывались. Сохранены все ограничения ежедневного датасета, в том числе условная формула ing_factor и неподписанная единица КВЧ.

## Что действительно есть по H2S

Повторно прочитаны все 128 книг и 228 листов в папках `УКК от ДО` и `УКК от ДО 2`. Явного столбца водного H2S в этих исходниках нет. Нефтяной H2S рассматривается как водный по пользовательскому правилу. В нефтяных колонках 550 числовых ячеек: 548 нулевых и 2 положительных. После удаления повторов это 275 разных записей: 274 нулевые и **одна положительная**.

Для направлений наших 66 труб найдены **30 исходных нефтяных ячеек**, соответствующие **15 уникальным сочетаниям УКК–дата–значение**. Все значения — **0 мг/м3**; они распространяются на 12 труб четырех направлений. При разнесении по трубам получаются 30 трубо-дат, из них 29 присутствуют в подготовленном календаре и 25 — в активном календаре нынешнего CSV. Поэтому формулировка «данных вовсе нет» была неточной: есть исходные нули. Действующий положительный фильтр H2S исключил их, и водная колонка осталась пустой. Лабораторное значение нуля (измеренный нуль, ниже предела определения либо иной код) из этих книг не установлено; считать его положительной концентрацией или автоматически заполнять им годы нельзя.

Единственная положительная проба: **3,7 мг/м3**, УКК **604**, **28.04.2026**, направление ПОТ **1750009316** (`ВЦТП - КНС-4 Ватинская`, простой участок `1751050477`, строка ПОТ 16240). Направление отсутствует среди целевых 66 труб. Источник: `УКК от ДО/апрель 2026/Шаблон данных по физ-хим составу жидкости апрель.xlsx`, `Sheet1`, **BI1784**. Та же проба повторяется в `УКК от ДО 2/2026/Шаблон данных по физ-хим составу жидкости апрель 2026.xlsx`, `Sheet1`, **BI1785**; второй раз не учитывается. По правилу мг/м3 -> мг/л это 0,0037, но подставлять его в другое направление оснований нет.

Примеры нулей целевых направлений: тот же апрельский файл, `Sheet1`, **BI934** (УКК 765, 06.04.2026) и **BI2164** (УКК 218, 30.04.2026). Полный список с файлами и ячейками — лист `H2S в исходниках` итоговой книги. В недельном датасете `H2S in Water Phase` остается пустым, поскольку средние считаются из готового ежедневного CSV.

## Проверка результата

Синтетические случаи совпали с pandas.groupby.mean: разные недели, неполные недели, NaN, нулевая дозировка, различный состав сегментов и несортированный исходник. Полный расчет проверил уникальность труба-дата-сегмент и сохранение числа исходных строк; для всех 47 усредняемых полей прошел контроль сумм с учетом числа наблюдений. Независимый проход по всему недельному CSV подтвердил схему, ключи, понедельничные даты, пропуски, границы средних и SHA-256. {len(verification['pandas_whole_week_checks'])} полных наборов недель дополнительно пересчитаны pandas из оригинальных дневных строк. Статус: **PASS**.
"""
    (HERE / "ИТОГ.md").write_text(text, encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
