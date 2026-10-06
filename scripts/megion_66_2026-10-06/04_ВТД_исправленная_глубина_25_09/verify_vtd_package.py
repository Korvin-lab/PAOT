from __future__ import annotations

import hashlib
import json
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
files = list((ROOT / 'input_vtd').glob('Мегион_ВТД_исправленная_глубина_*.xlsx'))
if len(files) != 1:
    raise RuntimeError(f'Ожидалась одна исправленная книга ВТД; найдено {len(files)}')
path = files[0]
digest = hashlib.sha256()
with path.open('rb') as src:
    for block in iter(lambda: src.read(8 * 1024 * 1024), b''):
        digest.update(block)
with zipfile.ZipFile(path) as archive:
    if archive.testzip() is not None:
        raise RuntimeError('Поврежден XLSX')
book = load_workbook(path, read_only=True, data_only=True)
if book.sheetnames != ['ВТД', 'Статистика']:
    raise RuntimeError('Некорректные листы ВТД')
rows = book['ВТД'].iter_rows(values_only=True)
header = next(rows)
if len(header) != 20 or header[10:12] != ('Толщ. стенки, мм', 'Остаточная толщина стенки, мм'):
    raise RuntimeError('Некорректная схема ВТД')
count = 0
cent = Decimal('0.01')
for row in rows:
    count += 1
    if len(row) != 20:
        raise RuntimeError(f'Число колонок в строке {count+1} не равно 20')
    wall, remain, depth = [Decimal(str(row[index])) for index in (10, 11, 16)]
    if not (wall > 0 and 0 <= remain < wall and 0 < depth <= 100):
        raise RuntimeError(f'Физически некорректная строка {count+1}')
    if any(value != value.quantize(cent) for value in (wall, remain, depth)):
        raise RuntimeError(f'Более двух десятичных знаков в строке {count+1}')
book.close()
report = {'status': 'PASS', 'vtd_rows': count, 'columns': 20, 'sha256': digest.hexdigest(), 'pe2_uses_vtd': False}
(ROOT / 'VTD_VALIDATION.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(report, ensure_ascii=False, indent=2))
