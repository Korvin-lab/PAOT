from __future__ import annotations

import json
from collections import Counter
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path('/Volumes/KINGSTON/Газпром работа/Прогнозирование утонения стенки трубы/Перевыгрузка по ВТД_Мегион_сентябрь2026')
BOOK = ROOT / 'ВТД_Мегион_исправленная_глубина_2026-09-25' / 'Мегион_ВТД_исправленная_глубина_2026_09_25.xlsx'
OUT = Path(__file__).resolve().parent


def cents(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


def main():
    wb = load_workbook(BOOK, read_only=True, data_only=True)
    assert wb.sheetnames == ['ВТД', 'Статистика']
    it = wb['ВТД'].iter_rows(values_only=True)
    header = next(it)
    assert len(header) == 20
    assert header[0] == 'Simple section ID'
    assert header[10:12] == ('Толщ. стенки, мм', 'Остаточная толщина стенки, мм')
    c = Counter()
    examples = []
    for row in it:
        c['rows'] += 1
        wall, residual, pct = [Decimal(str(row[i])) for i in (10, 11, 16)]
        assert wall == cents(wall) and residual == cents(residual) and pct == cents(pct)
        assert wall > 0 and 0 <= residual < wall and 0 < pct <= 100
        if row[0] is None or row[2] is None:
            c['missing_id_or_date'] += 1
        if str(row[0]) == '1751067170' and wall == Decimal('8') and residual == Decimal('7.4') and pct == Decimal('7.5'):
            c['fraction_example_pass'] += 1
        if str(row[0]) == '1750020578' and wall == Decimal('10') and residual == Decimal('9.92') and pct == Decimal('0.8'):
            c['early_subpercent_example_pass'] += 1
        if str(row[0]) == '1751045712' and wall == Decimal('10') and residual == Decimal('7.32'):
            c['measured_mm_example_pass'] += 1
        if c['rows'] <= 3:
            examples.append([row[0], row[2], str(wall), str(residual), str(pct)])
    wb.close()
    assert c['fraction_example_pass'] > 0
    assert c['early_subpercent_example_pass'] > 0
    assert c['measured_mm_example_pass'] > 0
    import zipfile
    with zipfile.ZipFile(BOOK) as archive:
        assert archive.testzip() is None
    report = {'status': 'PASS', **dict(c), 'columns': 20, 'excel_zip': 'PASS', 'examples': examples}
    (OUT/'output_validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
