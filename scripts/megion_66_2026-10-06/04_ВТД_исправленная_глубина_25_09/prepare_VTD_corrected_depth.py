from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import numpy as np
import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


OUTPUT_COLUMNS = [
    "Simple section ID", "ID освидетельствования", "Дата освидетельствования",
    "№ секции", "№ особ.", "Дист. от задвижки, м", "Отн. дист., м",
    "Длина трубы, м", "Описание особенности", "Комментарий", "Толщ. стенки, мм",
    "Остаточная толщина стенки, мм", "Угол, °", "Тип дефекта",
    "Длина особенности, мм", "Ширина особенности, мм", "Глубина дефекта, % WT(Dн)",
    "Широта, о", "Долгота, о", "Высота, м",
]

COLUMN_ALIASES = {
    "simple_id": ["id простого участка", "simple section id"],
    "inspection_id": ["id освидетельствования", "id обследования"],
    "inspection_date": ["дата освидетельствования", "дата обследования"],
    "section_no": ["№", "номер секции", "№ секции"],
    "feature_no": ["№ особенности", "номер особенности", "№ особ."],
    "address": ["адрес от начала простого", "дист. от задвижки, м"],
    "relative_distance": ["отн. дистанция,м", "отн. дистанция, м", "отн. дист., м"],
    "pipe_length": ["l простого", "длина трубы, м", "длина простого участка, м"],
    "description": ["идентификация особенности", "описание особенности"],
    "comment": ["комментарий"],
    "wall_thickness": ["толщина стенки элемента, мм", "толщ. стенки, мм", "номинальная толщина стенки, мм"],
    "angle": ["угловое положение дефекта", "угол, °"],
    "feature_type": ["тип особенности наим", "тип особенности", "тип дефекта", "идентификация особенности"],
    "feature_side": ["расположение дефекта", "местоположение дефекта", "поверхность", "сторона дефекта", "положение на стенке наим"],
    "feature_length": ["измеренная длина дефекта", "длина особенности, мм"],
    "feature_width": ["измеренная ширина дефекта", "ширина особенности, мм"],
    "depth_percent": ["максимальная измер.глубина,%", "максимальная измеренная глубина, % от толщины стенки", "глубина дефекта, % wt(dн)", "глубина дефекта, %"],
    "depth_mm": ["максимальная измер.глубина,мм", "максимальная измен. глубина, мм", "максимальная измеренная глубина, мм"],
}


def hundredths(value: float) -> float:
    return float(Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def normalized_depth_percent(raw_percent: float, raw_mm: float, wall: float, fraction_mode: bool) -> tuple[float, str]:
    if np.isfinite(raw_mm):
        if not np.isfinite(wall) or wall <= 0:
            return np.nan, 'measured_mm'
        return raw_mm / wall * 100, 'measured_mm'
    if not np.isfinite(raw_percent):
        return np.nan, 'missing'
    if 0 < raw_percent < 1 and fraction_mode:
        return raw_percent * 100, 'fraction'
    return raw_percent, 'percent'


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").strip()
    return re.sub(r"\s+", " ", text)


def norm(value: object) -> str:
    return re.sub(r"[^a-zа-я0-9%]+", " ", clean_text(value).lower().replace("ё", "е").replace("№", " номер ")).strip()


def find_column(headers: list[object], aliases: list[str]) -> int | None:
    normalized = [norm(value) for value in headers]
    for alias in aliases:
        target = norm(alias)
        for index, current in enumerate(normalized):
            if current == target:
                return index
    for alias in aliases:
        tokens = norm(alias).split()
        for index, current in enumerate(normalized):
            if tokens and all(token in current for token in tokens):
                return index
    return None


def parse_number(value: object) -> float:
    if value is None:
        return np.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value) if np.isfinite(value) else np.nan
    match = re.search(r"[-+]?\d+(?:[.,]\d+)?", clean_text(value).replace("−", "-"))
    return float(match.group(0).replace(",", ".")) if match else np.nan


def parse_angle(value: object) -> float:
    text = clean_text(value).replace(",", ".")
    clock = re.fullmatch(r"(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?", text)
    if clock:
        return round(int(clock.group(1)) * 15.0 + int(clock.group(2)) * 0.25 + int(clock.group(3) or 0) / 240.0, 4)
    return parse_number(value)


def identity(value: object) -> str:
    number = parse_number(value)
    if np.isfinite(number) and number.is_integer():
        return str(int(number))
    return clean_text(value)


def cell_date(book: xlrd.book.Book, cell: xlrd.sheet.Cell) -> str:
    if cell.ctype == xlrd.XL_CELL_DATE:
        return datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode)).date().isoformat()
    text = clean_text(cell.value)
    if not text:
        return ""
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], pattern).date().isoformat()
        except ValueError:
            pass
    return text


def main() -> None:
    root = Path(__file__).resolve().parent.parent
    output_dir = Path(__file__).resolve().parent
    output_xlsx = output_dir / "Мегион_ВТД_исправленная_глубина_2026_09_25.xlsx"
    report_path = output_dir / "ОТЧЕТ_СБОРКИ.json"

    input_files = sorted(path for path in root.glob("Выгрузка_ВТД_СН-МНГ_*.xls") if not path.name.startswith("._"))
    if not input_files:
        raise RuntimeError("Не найдены исходные выгрузки ВТД.")

    temporary_xlsx = output_xlsx.with_suffix(".tmp.xlsx")
    if temporary_xlsx.exists():
        temporary_xlsx.unlink()
    workbook = Workbook(write_only=True)
    data_sheet = workbook.create_sheet("ВТД")
    data_sheet.append(OUTPUT_COLUMNS)
    statistics: Counter[str] = Counter()
    source_rows: list[dict[str, object]] = []
    headers_reference: list[object] | None = None
    seen_measurement_pipe: set[tuple[str, str]] = set()

    for source_path in input_files:
        book = xlrd.open_workbook(str(source_path), on_demand=True, ignore_workbook_corruption=True)
        inspection_has_percent: set[str] = set()
        for sheet_name in book.sheet_names():
            scan = book.sheet_by_name(sheet_name)
            scan_headers = scan.row_values(0)
            scan_pct = find_column(scan_headers, COLUMN_ALIASES['depth_percent'])
            scan_id = find_column(scan_headers, COLUMN_ALIASES['inspection_id'])
            if scan_pct is None or scan_id is None:
                raise RuntimeError(f'{source_path.name}/{sheet_name}: отсутствует шкала глубины или ID освидетельствования')
            for scan_row in range(1, scan.nrows):
                raw = parse_number(scan.cell_value(scan_row, scan_pct))
                if np.isfinite(raw) and raw >= 1:
                    inspection_has_percent.add(identity(scan.cell_value(scan_row, scan_id)))
            book.unload_sheet(sheet_name)
        for sheet_name in book.sheet_names():
            sheet = book.sheet_by_name(sheet_name)
            headers = sheet.row_values(0)
            mapping = {key: find_column(headers, aliases) for key, aliases in COLUMN_ALIASES.items()}
            required = ["simple_id", "inspection_id", "inspection_date", "wall_thickness", "depth_percent", "depth_mm"]
            missing = [key for key in required if mapping[key] is None]
            if missing:
                raise RuntimeError(f"{source_path.name}/{sheet_name}: отсутствуют обязательные поля {missing}")
            rows_total = rows_depth = rows_thickness = rows_not_external = rows_output = 0
            for row_index in range(1, sheet.nrows):
                cells = sheet.row(row_index)
                if not any(clean_text(cell.value) for cell in cells):
                    continue
                rows_total += 1
                statistics["Исходных строк"] += 1
                value = lambda key: cells[mapping[key]].value if mapping.get(key) is not None else None
                source_percent = parse_number(value("depth_percent"))
                source_mm = parse_number(value("depth_mm"))
                thickness = parse_number(value("wall_thickness"))
                valid_thickness = np.isfinite(thickness) and thickness > 0
                fraction_mode = identity(value('inspection_id')) not in inspection_has_percent
                depth, depth_source = normalized_depth_percent(source_percent, source_mm, thickness, fraction_mode)
                valid_depth = np.isfinite(depth) and 0 < depth <= 100
                if valid_depth:
                    statistics[f"Глубина из {depth_source}"] += 1
                side_text = norm(value("feature_side")) if mapping.get("feature_side") is not None else ""
                valid_side = "внешн" not in side_text
                if valid_depth:
                    rows_depth += 1
                    statistics["Валидная глубина > 0 и <= 100%"] += 1
                if valid_thickness:
                    rows_thickness += 1
                    statistics["Положительная номинальная толщина"] += 1
                if valid_side:
                    rows_not_external += 1
                    statistics["Не внешнее расположение"] += 1
                if not (valid_depth and valid_thickness and valid_side):
                    continue

                simple_id = identity(value("simple_id"))
                inspection_id = identity(value("inspection_id"))
                measure_col = find_column(headers, ["id замера"])
                measure_id = identity(cells[measure_col].value) if measure_col is not None else ""
                duplicate_key = (measure_id, simple_id)
                if measure_id and simple_id and duplicate_key in seen_measurement_pipe:
                    statistics["Полные дубли ID замера + простой участок"] += 1
                    continue
                if measure_id and simple_id:
                    seen_measurement_pipe.add(duplicate_key)

                wall_output = hundredths(thickness)
                residual = hundredths(thickness - (source_mm if depth_source == 'measured_mm' else thickness * depth / 100))
                if wall_output <= 0 or residual < 0 or residual >= wall_output:
                    statistics['Отклонено после округления до 0.01 мм'] += 1
                    continue
                depth_output = hundredths(depth)
                if depth_output <= 0:
                    statistics['Отклонено: глубина исчезла после округления до 0.01%'] += 1
                    continue
                data_sheet.append([
                    simple_id, inspection_id, cell_date(book, cells[mapping["inspection_date"]]),
                    value("section_no"), value("feature_no"), parse_number(value("address")),
                    parse_number(value("relative_distance")), parse_number(value("pipe_length")),
                    clean_text(value("description")), clean_text(value("comment")), wall_output, residual,
                    parse_angle(value("angle")), clean_text(value("feature_type")),
                    parse_number(value("feature_length")), parse_number(value("feature_width")), depth_output,
                    "", "", "",
                ])
                rows_output += 1
                statistics["Итоговых строк"] += 1
            source_rows.append({
                "Файл": source_path.name, "Лист": sheet_name, "Исходных строк": rows_total,
                "Валидная глубина": rows_depth, "Положительная толщина": rows_thickness,
                "Не внешние": rows_not_external, "Выгружено": rows_output,
                "Заголовки совпадают": headers_reference is None or headers == headers_reference,
            })
            headers_reference = headers if headers_reference is None else headers_reference
            book.unload_sheet(sheet_name)
        book.release_resources()

    statistics["Листов"] = len(source_rows)
    statistics["Файлов"] = len(input_files)
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = f"A1:T{statistics['Итоговых строк'] + 1}"
    widths = [18, 22, 18, 12, 12, 20, 16, 16, 30, 30, 18, 25, 12, 30, 24, 24, 27, 15, 15, 15]
    for index, width in enumerate(widths, 1):
        data_sheet.column_dimensions[get_column_letter(index)].width = width

    stats_sheet = workbook.create_sheet("Статистика")
    stats_sheet.append(["Показатель", "Значение"])
    stats_sheet.append(["Правило типа дефекта", "Не фильтруется: коррозия/потеря металла не требуется"])
    stats_sheet.append(["Правило расположения", "Исключены только внешние особенности"])
    stats_sheet.append(["Правило глубины", "Строго > 0 и <= 100%"])
    stats_sheet.append(["Источник глубины", "Измеренные мм имеют приоритет; иначе шкала %/доля определяется по ID освидетельствования"])
    stats_sheet.append(["Точность", "Номинальная и остаточная толщина, а также глубина в % округлены до 0.01"])
    stats_sheet.append(["Правило номинальной толщины", "Строго > 0 мм"])
    for key in ["Файлов", "Листов", "Исходных строк", "Валидная глубина > 0 и <= 100%", "Положительная номинальная толщина", "Не внешнее расположение", "Полные дубли ID замера + простой участок", "Итоговых строк"]:
        stats_sheet.append([key, int(statistics[key])])
    stats_sheet.append([])
    stats_sheet.append(["Файл", "Лист", "Исходных строк", "Валидная глубина", "Положительная толщина", "Не внешние", "Выгружено", "Заголовки совпадают"])
    for item in source_rows:
        stats_sheet.append([item[key] for key in ["Файл", "Лист", "Исходных строк", "Валидная глубина", "Положительная толщина", "Не внешние", "Выгружено", "Заголовки совпадают"]])
    stats_sheet.freeze_panes = "A2"
    stats_sheet.column_dimensions["A"].width = 45
    stats_sheet.column_dimensions["B"].width = 55
    for column in range(3, 9):
        stats_sheet.column_dimensions[get_column_letter(column)].width = 22

    workbook.save(temporary_xlsx)
    temporary_xlsx.replace(output_xlsx)

    # Independent read-back: schema, exact row count, and retained validity constraints.
    check_book = load_workbook(output_xlsx, read_only=True, data_only=True)
    check_sheet = check_book["ВТД"]
    check_header = [cell.value for cell in next(check_sheet.iter_rows(min_row=1, max_row=1))]
    checked_rows = 0
    invalid_rows = 0
    for row in check_sheet.iter_rows(min_row=2, values_only=True):
        checked_rows += 1
        thickness, residual, depth = row[10], row[11], row[16]
        if not (isinstance(thickness, (int, float)) and thickness > 0 and isinstance(depth, (int, float)) and 0 < depth <= 100 and isinstance(residual, (int, float)) and 0 <= residual < thickness):
            invalid_rows += 1
    check_book.close()
    if check_header != OUTPUT_COLUMNS or checked_rows != statistics["Итоговых строк"] or invalid_rows:
        raise RuntimeError("Независимая проверка итоговой книги не пройдена.")

    report = {
        "статус": "PASS",
        "входные_файлы": [path.name for path in input_files],
        "правило": {
            "потеря_металла": "не фильтруется",
            "внешняя_сторона": "исключена",
            "глубина": "мм при наличии; иначе доля при 0<x<1 лишь у освидетельствований без значений >=1, проценты в остальных; итог 0<x<=100%",
            "номинальная_толщина": "> 0 мм",
        },
        "статистика": dict(statistics),
        "по_листам": source_rows,
        "проверка_выхода": {"колонки": check_header, "строк": checked_rows, "некорректных_строк": invalid_rows},
        "файл": str(output_xlsx),
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
