import json,re,zipfile,xml.etree.ElementTree as ET
from collections import defaultdict,Counter
from pathlib import Path
from datetime import datetime,date
from openpyxl import load_workbook
import numpy as np
ROOT=Path('/Volumes/KINGSTON/Газпром работа/Прогнозирование утонения стенки трубы')
WORK=ROOT/'работа с json'; OUT=Path(__file__).resolve().parent
PACKAGE=WORK/'MEGION_WINDOWS_FULL_PE2__PREPARING_2026-09-18'
cache=json.loads((PACKAGE/'audit_megion/ukk_value_extract_cache.json').read_text())
active=set()
import pandas as pd
active=set(pd.read_csv(PACKAGE/'input/chem_daily_megion.csv',usecols=['id простого участка'])['id простого участка'].astype(str))

def norm(x):
 s=str(x or '').strip();return s[:-2] if s.endswith('.0') and s[:-2].isdigit() else s

def ukk(x):return re.sub(r'[^0-9A-ZА-Я]+','',str(x or '').upper().replace('Ё','Е')).replace('УКК','').replace('N','').replace('№','')
def day(x):
 if isinstance(x,datetime):return x.date().isoformat()
 if isinstance(x,date):return x.isoformat()
 for fmt in ('%d.%m.%Y','%Y-%m-%d','%d.%m.%Y %H:%M:%S'):
  try:return datetime.strptime(str(x),fmt).date().isoformat()
  except:pass
 return None
def num(x):
 try:return float(str(x).replace('\xa0','').replace(' ','').replace(',','.'))
 except:return None
p=load_workbook(ROOT/'Перечень_ответственных_трубопроводов_СН_МНГ 2.xlsx',read_only=True,data_only=True)
by_ukk=defaultdict(set); target_dirs=defaultdict(set)
for row in p.active.iter_rows(min_row=3,values_only=True):
 pid=norm(row[48] if len(row)>48 else None);did=norm(row[42] if len(row)>42 else None);uid=ukk(row[244] if len(row)>244 else None)
 if pid in active: target_dirs[did].add(pid)
 if uid and did:by_ukk[uid].add(did)
p.close()
valid=defaultdict(list)
for z in cache['structures']:
 if z['Валидный физ-хим шаблон']:valid[z['Файл']].append(z['Лист'])
allrows=[]; mapped=[]; ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
for f, sheets in valid.items():
 wb=load_workbook(ROOT/f,read_only=True,data_only=True)
 with zipfile.ZipFile(ROOT/f) as archive:
  for sn in sheets:
   ws=wb[sn];positions=set()
   with archive.open(ws._worksheet_path.lstrip('/')) as stream:
    for _,el in ET.iterparse(stream,events=('end',)):
     if el.tag!=ns+'row':continue
     rn=int(el.get('r'))
     if rn>=4:
      for c in el:
       if c.tag==ns+'c' and c.get('r')==f'BD{rn}':
        v=c.find(ns+'v')
        if v is not None and num(v.text) is not None and num(v.text)>0:positions.add(rn)
        break
     el.clear()
   if not positions:continue
   head=[(ws.cell(r,56).value) for r in range(1,5)]
   uid='';dt=None
   for rn,row in enumerate(ws.iter_rows(min_row=4,max_row=max(positions),min_col=14,max_col=56,values_only=True),4):
    if row[0] is not None and str(row[0]).strip():uid=ukk(row[0])
    d=day(row[6])
    if d:dt=d
    if rn not in positions:continue
    v=num(row[42]);
    if not v or v<=0:continue
    item={'file':f,'sheet':sn,'row':rn,'ukk':uid,'date':dt,'raw':v,'header':str(head)}
    allrows.append(item)
    for did in by_ukk.get(uid,()):
     if did in target_dirs and dt:
      mapped.append({**item,'direction':did,'pipes':sorted(target_dirs[did])})
 wb.close()
print('all positive rows',len(allrows),'mapped direction rows',len(mapped),'directions',len(set(x['direction'] for x in mapped)),'pipes',len(set(p for x in mapped for p in x['pipes'])))
for name,rows in [('all',allrows),('mapped',mapped)]:
 vals=np.array([x['raw'] for x in rows]);print(name,'min/quantiles/max',np.percentile(vals,[0,1,5,25,50,75,95,99,100]).tolist(),'lt100',int((vals<100).sum()),'100-1000',int(((vals>=100)&(vals<1000)).sum()),'>=1000',int((vals>=1000).sum()))
 byfile=defaultdict(list)
 for x in rows:byfile[x['file']].append(x['raw'])
 print('files',name,len(byfile))
 for f,v in sorted(byfile.items()):
  a=np.array(v);print(f,len(a),'min-med-max',float(a.min()),float(np.median(a)),float(a.max()))
(OUT/'ukk_mineralization_all_positive.json').write_text(json.dumps(allrows,ensure_ascii=False))
(OUT/'ukk_mineralization_mapped.json').write_text(json.dumps(mapped,ensure_ascii=False))
