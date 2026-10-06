#!/usr/bin/env python3
"""Prepare date-specific chemistry replacements without rerunning PE2."""

from __future__ import annotations

import importlib.util
import json
import math
import re
from collections import defaultdict
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path('/Volumes/KINGSTON/Газпром работа/Прогнозирование утонения стенки трубы')
WORK = ROOT / 'работа с json'
PACKAGE = WORK / 'MEGION_WINDOWS_FULL_PE2__PREPARING_2026-09-18'
OUT = Path(__file__).resolve().parent
FIELDS = ('CO2', 'pH', 'H2S in Gas Phase')


def old_builder():
    path = PACKAGE / 'build_megion_windows_inputs.py'
    spec = importlib.util.spec_from_file_location('megion_inputs', path)
    module = importlib.util.module_from_spec(spec)
    assert spec and spec.loader
    spec.loader.exec_module(module)
    return module


def build_direct(module, days, pot):
    by_ukk = defaultdict(set)
    book = load_workbook(ROOT / 'Перечень_ответственных_трубопроводов_СН_МНГ 2.xlsx',
                         read_only=True, data_only=True)
    for row in book.active.iter_rows(min_row=3, values_only=True):
        did = module.norm_id(row[42] if len(row) > 42 else None)
        ukk = re.sub(r'[^0-9A-ZА-Я]+', '', module.text(row[244] if len(row) > 244 else '').upper().replace('Ё', 'Е')).replace('УКК', '').replace('N', '').replace('№', '')
        if ukk and did:
            by_ukk[ukk].add(did)
    book.close()
    target_by_dir = defaultdict(list)
    for pid in days.id.unique():
        target_by_dir[pot[pid]['direction_id']].append(pid)
    records = json.loads((PACKAGE / 'audit_megion/ukk_value_extract_cache.json').read_text(encoding='utf-8'))['records']
    rows = []
    for rec in records:
        for did in by_ukk.get(module.text(rec['Номер УКК']), ()):
            for pid in target_by_dir.get(did, ()):
                rows.append((pid, rec['Дата'], rec['CO2 в водной фазе'], rec['pH'], rec['H2S в газовой фазе']))
    raw = pd.DataFrame(rows, columns=['id', 'date', *FIELDS])
    direct = raw.groupby(['id', 'date'], as_index=False).median(numeric_only=True)
    direct['date'] = pd.to_datetime(direct['date'])
    return days.merge(direct, on=['id', 'date'], how='left'), len(records), len(raw)


def fill_donor_series(module, frame, field):
    frame = frame.copy().sort_values(['id', 'date']).reset_index(drop=True)
    frame[field] = pd.to_numeric(frame[field], errors='coerce').where(lambda x: x > 0)
    donor_ids = set(frame.loc[frame[field].notna(), 'id'])
    filled = []
    for pid in sorted(donor_ids):
        part = frame[frame.id.eq(pid)].copy()
        values, _ = module.fill_one_series(part[field], part['date'], math.nan, lambda s: s > 0)
        part[field] = values
        assert part[field].notna().all() and (part[field] > 0).all()
        filled.append(part[['id', 'date', field]])
    if not filled:
        raise AssertionError(f'No donor pipes for {field}')
    return pd.concat(filled, ignore_index=True), donor_ids


def daily_median(donors, all_dates, field):
    observed = donors.groupby('date')[field].median().sort_index()
    full_index = pd.DatetimeIndex(sorted(set(all_dates).union(observed.index)))
    complete = observed.reindex(full_index).interpolate(method='time', limit_area='inside').bfill().ffill()
    result = complete.reindex(pd.DatetimeIndex(all_dates))
    assert result.notna().all() and (result > 0).all()
    return result.to_numpy(float), len(observed)


def main():
    OUT.mkdir(exist_ok=True)
    module = old_builder()
    old_chem = pd.read_csv(PACKAGE / 'input/chem_daily_megion.csv',
                           dtype={'id простого участка': str, 'Дата контроля': str})
    old_gas = pd.read_csv(PACKAGE / 'input/h2s_two_phase_daily_base.csv', dtype={'id': str, 'date': str})
    assert len(old_chem) == len(old_gas) == 99212
    days = old_chem[['id простого участка', 'Дата контроля']].rename(
        columns={'id простого участка': 'id', 'Дата контроля': 'date'})
    days['date'] = pd.to_datetime(days['date'], format='%Y-%m-%d')
    assert not days.duplicated(['id', 'date']).any() and days.id.nunique() == 66
    pot = module.load_pot()
    direct, records, raw_rows = build_direct(module, days, pot)
    old = days.copy()
    old['CO2'] = old_chem['CO2'].to_numpy(float)
    old['pH'] = old_chem['pH'].to_numpy(float)
    old['H2S in Gas Phase'] = old_gas['H2S in Gas Phase'].to_numpy(float)
    if not old_gas[['id','date']].assign(date=lambda x: pd.to_datetime(x.date)).equals(days):
        raise AssertionError('Prepared H2S gas series not aligned with chemistry dates')
    replacement = days.copy()
    report = {'source_records': records, 'mapped_pipe_date_records': raw_rows, 'fields': {}}
    for field in FIELDS:
        donors, donor_ids = fill_donor_series(module, direct, field)
        old_donor = old[old.id.isin(donor_ids)].sort_values(['id', 'date']).reset_index(drop=True)
        new_donor = donors.sort_values(['id', 'date']).reset_index(drop=True)
        assert old_donor[['id', 'date']].equals(new_donor[['id', 'date']])
        difference = np.abs(old_donor[field].to_numpy(float) - new_donor[field].to_numpy(float))
        if not np.allclose(old_donor[field].to_numpy(float), new_donor[field].to_numpy(float), rtol=1e-12, atol=1e-10):
            raise AssertionError(f'Donor series differs from original prepared input: {field}; max={difference.max()}')
        median, n_obs_days = daily_median(donors, days['date'], field)
        replacement[field] = old[field].where(old.id.isin(donor_ids), median)
        affected = ~old.id.isin(donor_ids)
        report['fields'][field] = {
            'donor_pipes': len(donor_ids), 'replacement_pipes': int(old.loc[affected, 'id'].nunique()),
            'replacement_pipe_dates': int(affected.sum()), 'days_with_donor_value': n_obs_days,
            'donor_max_absolute_difference': float(difference.max()),
            'daily_median_min': float(np.min(median)), 'daily_median_max': float(np.max(median)),
            'replacement_unique_values': int(replacement.loc[affected, field].nunique()),
        }
    replacement['pCO2'] = replacement['CO2'] / 4410.0
    assert replacement[list(FIELDS)].notna().all().all()
    assert replacement[list(FIELDS)].gt(0).all().all()
    replacement['date'] = replacement['date'].dt.strftime('%Y-%m-%d')
    replacement = replacement.sort_values(['id', 'date'])
    replacement.to_csv(OUT / 'daily_chemistry_corrected.csv', index=False, encoding='utf-8')
    # Only changed ID-dates need rewriting in the 95-GB segment file.
    old_ordered = old.assign(date=old.date.dt.strftime('%Y-%m-%d')).sort_values(['id', 'date'])
    changed = replacement.set_index(['id', 'date']).join(
        old_ordered.set_index(['id', 'date'])[list(FIELDS)], rsuffix='_old')
    mask = np.zeros(len(changed), dtype=bool)
    for field in FIELDS:
        mask |= ~np.isclose(changed[field].to_numpy(float), changed[f'{field}_old'].to_numpy(float),
                            rtol=1e-12, atol=1e-10)
    patch = changed.loc[mask, [*FIELDS, 'pCO2']].reset_index()
    patch.to_csv(OUT / 'daily_replacements.csv', index=False, encoding='utf-8')
    report['changed_id_dates'] = len(patch)
    report['changed_pipes'] = patch.id.nunique()
    report['unchanged_mineralization_mg_l'] = float(old_chem['Общая минерализация'].iloc[0])
    (OUT / 'DAILY_PATCH_REPORT.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
